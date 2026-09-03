from __future__ import annotations

import csv
import tempfile
import unittest
from pathlib import Path
from unittest.mock import patch

from openpyxl import Workbook
from pptx import Presentation

from pocketrag_backend.documents import DocumentExtractor


SETTINGS = {
    "textSplitter": "token",
    "chunkSize": 300,
    "chunkOverlap": 40,
    "minimumChunkSize": 0,
    "enrichedMetadata": True,
    "preserveHeadings": True,
}


class DocumentExtractorTests(unittest.TestCase):
    def setUp(self) -> None:
        self.temporary = tempfile.TemporaryDirectory()
        self.root = Path(self.temporary.name)
        self.extractor = DocumentExtractor()

    def tearDown(self) -> None:
        self.temporary.cleanup()

    def test_markdown_citations_keep_section_and_lines(self) -> None:
        path = self.root / "notes.md"
        path.write_text("# Overview\nFirst section.\n# Details\nSecond section.\n")

        chunks = self.extractor.extract(
            str(path), path.name, {**SETTINGS, "minimumChunkSize": 200}
        )

        self.assertEqual([chunk.heading for chunk in chunks], ["Overview", "Details"])
        self.assertEqual(chunks[0].location, "Overview, Lines 1-2")
        self.assertEqual(chunks[1].location, "Details, Lines 3-4")

    def test_csv_citations_use_source_row_ranges(self) -> None:
        path = self.root / "records.csv"
        with path.open("w", newline="") as handle:
            csv.writer(handle).writerows([["Name", "Value"], ["One", 1], ["Two", 2]])

        chunks = self.extractor.extract(str(path), path.name, SETTINGS)

        self.assertEqual(len(chunks), 1)
        self.assertEqual(chunks[0].location, "Rows 2-3")

    def test_excel_overlap_never_crosses_sheet_boundaries(self) -> None:
        path = self.root / "workbook.xlsx"
        workbook = Workbook()
        first = workbook.active
        first.title = "North"
        first.append(["Region", "Fact"])
        first.append(["North", "NORTH_ONLY_TEXT"])
        second = workbook.create_sheet("South")
        second.append(["Region", "Fact"])
        second.append(["South", "SOUTH_ONLY_TEXT"])
        workbook.save(path)

        chunks = self.extractor.extract(
            str(path), path.name, {**SETTINGS, "minimumChunkSize": 200}
        )

        self.assertEqual([chunk.sheet_name for chunk in chunks], ["North", "South"])
        self.assertNotIn("NORTH_ONLY_TEXT", chunks[1].text)
        self.assertEqual(chunks[1].location, "Sheet South, A2:B2")

    def test_excel_cells_preserve_typed_values_headers_and_units(self) -> None:
        path = self.root / "pricing.xlsx"
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "Price Estimate"
        sheet.append(["Item", "Unit Price USD", "Lead Time"])
        sheet.append(["H200", 7832, "14 days"])
        workbook.save(path)

        cells = self.extractor.extract_spreadsheet_cells(str(path))

        price = next(cell for cell in cells if cell.column_name == "Unit Price USD")
        lead = next(cell for cell in cells if cell.column_name == "Lead Time")
        self.assertEqual(price.numeric_value, 7832.0)
        self.assertEqual(price.unit, "USD")
        self.assertEqual(price.cell_address, "B2")
        self.assertEqual(lead.numeric_value, 14.0)
        self.assertEqual(lead.unit, "days")

    def test_excel_display_values_keep_decimal_precision_without_scientific_notation(self) -> None:
        path = self.root / "totals.xlsx"
        workbook = Workbook()
        sheet = workbook.active
        sheet.append(["Metric", "Value"])
        sheet.append(["Product Total", 3074946.94])
        sheet.append(["HBM occupancy", 0.7620391706268562])
        workbook.save(path)

        cells = self.extractor.extract_spreadsheet_cells(str(path))

        values = {cell.cell_address: cell.display_value for cell in cells}
        self.assertEqual(values["B2"], "3074946.94")
        self.assertEqual(values["B3"], "0.7620391706268562")

        chunks = self.extractor.extract(str(path), path.name, SETTINGS)
        self.assertIn("| Product Total | 3074946.94 |", chunks[0].text)

    def test_excel_detects_header_below_title_rows(self) -> None:
        path = self.root / "boq.xlsx"
        workbook = Workbook()
        sheet = workbook.active
        sheet.append(["Bank Indonesia AI Infrastructure Proposal"])
        sheet.append([""])
        sheet.append(["Part Number", "Description", "Quantity", "Unit Price USD"])
        sheet.append(["UCSC-C240-M8", "Compute node", 4, 12000])
        workbook.save(path)

        cells = self.extractor.extract_spreadsheet_cells(str(path))
        chunks = self.extractor.extract(str(path), path.name, SETTINGS)

        self.assertEqual({cell.column_name for cell in cells}, {
            "Part Number", "Description", "Quantity", "Unit Price USD"
        })
        self.assertIn("UCSC-C240-M8", chunks[0].text)
        self.assertEqual(chunks[0].cell_range, "A4:D4")

    def test_excel_does_not_mistake_a_wide_data_row_for_the_header(self) -> None:
        path = self.root / "proposal.xlsx"
        workbook = Workbook()
        sheet = workbook.active
        sheet.append([""])
        sheet.append(["Customer proposal"])
        sheet.append([""])
        sheet.append(["No", "Product", "Product Description", "Qty", "Cisco propose:", "Cisco note"])
        sheet.append(
            [
                1,
                "NX-8150G-G10-6517P-CM",
                "NX-8150G-G10, 1 Node; 2x Intel Xeon processors per node",
                1,
                "Intel I6724P",
                "",
            ]
        )
        sheet.append([2, "C-GPU-L40S", "NVIDIA L40S GPU Card", 1, "", ""])
        workbook.save(path)

        chunks = self.extractor.extract(str(path), path.name, SETTINGS)

        self.assertIn("NX-8150G-G10-6517P-CM", chunks[0].text)
        self.assertIn("NVIDIA L40S GPU Card", chunks[0].text)
        self.assertEqual(chunks[0].cell_range, "A5:F6")

    def test_excel_combines_secondary_headers_without_indexing_them_as_data(self) -> None:
        path = self.root / "boq-secondary.xlsx"
        workbook = Workbook()
        sheet = workbook.active
        sheet.append([""])
        sheet.append(["", "No", "Product", "Description", "Qty"])
        sheet.append(["", "Project title", "", "", "", "Cisco propose", "Cisco note"])
        sheet.append(["", 1, "SKU-1", "Server", 2, "C220", "Validated"])
        workbook.save(path)

        cells = self.extractor.extract_spreadsheet_cells(str(path))
        chunks = self.extractor.extract(str(path), path.name, SETTINGS)

        row = [cell for cell in cells if cell.row_index == 4]
        self.assertIn("Cisco propose", {cell.column_name for cell in row})
        self.assertIn("Cisco note", {cell.column_name for cell in row})
        self.assertNotIn("Project title", {cell.display_value for cell in cells})
        self.assertEqual(chunks[0].cell_range, "A4:G4")

    def test_excel_preserves_single_cell_subheader_in_composite_column_name(self) -> None:
        path = self.root / "requirements.xlsx"
        workbook = Workbook()
        sheet = workbook.active
        sheet.append(["No", "Category", "Minimum Requirement", "CISCO"])
        sheet.append(["", "", "", "UCS C220 M7"])
        sheet.append([1, "Disk", "500GB RAID1", "960GB RAID1"])
        workbook.save(path)

        cells = self.extractor.extract_spreadsheet_cells(str(path))

        proposal = next(cell for cell in cells if cell.cell_address == "D3")
        self.assertEqual(proposal.column_name, "CISCO > UCS C220 M7")

    def test_excel_key_value_sheet_uses_metric_and_value_headers(self) -> None:
        path = self.root / "sizing.xlsx"
        workbook = Workbook()
        sheet = workbook.active
        sheet.append(["Sizing model"])
        sheet.append(["WORKLOAD PARAMETERS"])
        sheet.append(["Model", "Qwen"])
        sheet.append(["Instances", 6])
        sheet.append(["Input tokens", 5000])
        sheet.append(["Output tokens", 3000])
        sheet.append(["Runtime reserve (GiB)", 20])
        sheet.append(["HBM occupancy", 0.75])
        sheet.append(["Total GPUs", 14])
        workbook.save(path)

        cells = self.extractor.extract_spreadsheet_cells(str(path))
        chunks = self.extractor.extract(str(path), path.name, SETTINGS)

        total = next(cell for cell in cells if cell.display_value == "Total GPUs")
        value = next(cell for cell in cells if cell.cell_address == "B9")
        self.assertEqual(total.column_name, "Metric")
        self.assertEqual(value.column_name, "Value")
        self.assertIn("| Metric | Value |", chunks[0].text)

    def test_document_metadata_builds_local_section_map(self) -> None:
        path = self.root / "notes.md"
        path.write_text("# Architecture\nCisco ACI design guidance.\n# Operations\nRunbook details.\n")
        chunks = self.extractor.extract(str(path), path.name, SETTINGS)

        metadata = self.extractor.document_metadata(path.name, chunks)

        self.assertEqual(metadata["title"], "notes")
        self.assertEqual(metadata["headings"], ["Architecture", "Operations"])
        self.assertEqual(len(metadata["sections"]), 2)

    def test_pdf_text_is_extracted_locally_with_page_citations(self) -> None:
        class FakeTextPage:
            def get_text_range(self) -> str:
                return "Local PDF content"

            def close(self) -> None:
                pass

        class FakePage:
            def get_textpage(self) -> FakeTextPage:
                return FakeTextPage()

            def close(self) -> None:
                pass

        class FakeDocument:
            def __len__(self) -> int:
                return 1

            def __getitem__(self, _index: int) -> FakePage:
                return FakePage()

            def close(self) -> None:
                pass

        path = self.root / "local.pdf"
        path.write_bytes(b"pdf")
        with patch("pocketrag_backend.documents.pdfium.PdfDocument", return_value=FakeDocument()):
            chunks = self.extractor.extract(str(path), path.name, SETTINGS)

        self.assertEqual(chunks[0].text, "Local PDF content")
        self.assertEqual(chunks[0].location, "Page 1")

    def test_powerpoint_never_merges_text_across_slide_boundaries(self) -> None:
        path = self.root / "design.pptx"
        presentation = Presentation()
        first = presentation.slides.add_slide(presentation.slide_layouts[1])
        first.shapes.title.text = "Load Balancer Design"
        first.placeholders[1].text = "Use SNAT for return traffic."
        second = presentation.slides.add_slide(presentation.slide_layouts[1])
        second.shapes.title.text = "Firewall Design"
        second.placeholders[1].text = "Use VzAny with PBR."
        presentation.save(path)

        chunks = self.extractor.extract(
            str(path), path.name, {**SETTINGS, "minimumChunkSize": 300}
        )

        self.assertEqual([chunk.location for chunk in chunks], ["Slide 1", "Slide 2"])
        self.assertEqual(
            [chunk.heading for chunk in chunks],
            ["Load Balancer Design", "Firewall Design"],
        )
        self.assertIn("SNAT", chunks[0].text)
        self.assertNotIn("VzAny", chunks[0].text)
        self.assertIn("VzAny", chunks[1].text)
        self.assertNotIn("SNAT", chunks[1].text)

    def test_token_chunks_break_on_source_lines_before_raw_token_boundaries(self) -> None:
        text = "\n".join(
            [
                "First complete source sentence about Intersight operations.",
                "Second complete source sentence about Cisco Validated Designs.",
                "Third complete source sentence about air cooling limitations.",
            ]
        )
        parts = self.extractor._split_text(text, 12, "token")

        self.assertGreater(len(parts), 1)
        self.assertTrue(all(part[0].isupper() for part in parts))
        self.assertTrue(all(part.endswith(".") for part in parts))

    def test_scanned_pdf_uses_local_ocr_with_page_citations(self) -> None:
        class FakeTextPage:
            def get_text_range(self) -> str:
                return ""

            def close(self) -> None:
                pass

        class FakeBitmap:
            def to_pil(self) -> object:
                return object()

            def close(self) -> None:
                pass

        class FakePage:
            def get_textpage(self) -> FakeTextPage:
                return FakeTextPage()

            def render(self, scale: int) -> FakeBitmap:
                self.scale = scale
                return FakeBitmap()

            def close(self) -> None:
                pass

        class FakeDocument:
            def __len__(self) -> int:
                return 1

            def __getitem__(self, _index: int) -> FakePage:
                return FakePage()

            def close(self) -> None:
                pass

        class FakeResult:
            txts = ["Recognized local text"]

        path = self.root / "scanned.pdf"
        path.write_bytes(b"pdf")
        with (
            patch("pocketrag_backend.documents.pdfium.PdfDocument", return_value=FakeDocument()),
            patch("rapidocr.RapidOCR", return_value=lambda _image: FakeResult()),
        ):
            chunks = self.extractor.extract(str(path), path.name, {**SETTINGS, "pdfOcr": True})

        self.assertEqual(chunks[0].text, "Recognized local text")
        self.assertEqual(chunks[0].location, "Page 1")

    def test_standalone_image_uses_local_ocr(self) -> None:
        class FakeImage:
            def __enter__(self) -> "FakeImage":
                return self

            def __exit__(self, *_args: object) -> None:
                pass

            def convert(self, _mode: str) -> object:
                return object()

        class FakeResult:
            txts = ["Cisco NetApp cue card", "Key customer outcome"]

        path = self.root / "cue-card.png"
        path.write_bytes(b"image")
        with (
            patch("PIL.Image.open", return_value=FakeImage()),
            patch("rapidocr.RapidOCR", return_value=lambda _image: FakeResult()),
        ):
            chunks = self.extractor.extract(str(path), path.name, SETTINGS)

        self.assertIn("Cisco NetApp cue card", chunks[0].text)
        self.assertEqual(chunks[0].location, "Image text")


if __name__ == "__main__":
    unittest.main()
