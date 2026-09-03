from __future__ import annotations

import csv
from datetime import date, datetime
from itertools import chain
import math
import re
from dataclasses import dataclass, field
from pathlib import Path
from typing import Any, Iterable

import tiktoken
import pypdfium2 as pdfium
from docling.chunking import HybridChunker
from docling.datamodel.base_models import InputFormat
from docling.datamodel.pipeline_options import PdfPipelineOptions
from docling.document_converter import DocumentConverter, PdfFormatOption
from docling_core.transforms.chunker.tokenizer.openai import OpenAITokenizer
from openpyxl import load_workbook
from openpyxl.utils import get_column_letter
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE


@dataclass
class DocumentChunk:
    text: str
    embedding_text: str = ""
    heading: str = ""
    page_start: int = 0
    page_end: int = 0
    paragraph_start: int = 0
    paragraph_end: int = 0
    sheet_name: str = ""
    cell_range: str = ""
    row_start: int = 0
    row_end: int = 0
    line_start: int = 0
    line_end: int = 0
    location: str = ""
    metadata: dict[str, Any] = field(default_factory=dict)


@dataclass
class SpreadsheetCell:
    sheet_name: str
    row_index: int
    column_index: int
    column_name: str
    cell_address: str
    display_value: str
    raw_value: str
    value_type: str
    numeric_value: float | None = None
    unit: str = ""
    formula: str = ""


def _display_value(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, (datetime, date)):
        return value.isoformat()
    if isinstance(value, float) and math.isfinite(value):
        # Excel formula caches sometimes expose harmless IEEE-754 tails on large
        # currency totals (for example 3074946.9400000004). Fifteen significant
        # digits removes that artifact while small ratios retain full precision.
        text = format(value, ".15g") if abs(value) >= 1000 else repr(value)
        return text[:-2] if text.endswith(".0") else text
    return str(value)


def _numeric_value(value: Any) -> float | None:
    if isinstance(value, bool):
        return None
    if isinstance(value, (int, float)) and math.isfinite(float(value)):
        return float(value)
    if not isinstance(value, str):
        return None
    match = re.search(r"(?:[$€£]\s*)?(-?\d[\d,]*(?:\.\d+)?)", value)
    if not match:
        return None
    try:
        return float(match.group(1).replace(",", ""))
    except ValueError:
        return None


def _unit(header: str, value: str) -> str:
    candidate = f"{header} {value}".casefold()
    patterns = (
        (r"[$]|\busd\b", "USD"),
        (r"\bidr\b|\brp\.?\b", "IDR"),
        (r"\b(days?|weeks?|months?|years?|hours?|minutes?|seconds?)\b", None),
        (r"\b(gb|tb|mb|kb)\b", None),
        (r"\b(k?w|mw|gw)\b", None),
        (r"%", "%"),
    )
    for pattern, fixed in patterns:
        match = re.search(pattern, candidate, re.IGNORECASE)
        if match:
            return fixed or match.group(1).lower()
    return ""


def _header_row_position(rows: list[tuple[int, list[str]]]) -> int:
    """Choose the most table-like row near the top instead of assuming row one."""
    if not rows:
        return 0
    best_position = 0
    best_score = float("-inf")
    for position, (_row_index, values) in enumerate(rows[:25]):
        populated = [value.strip() for value in values if value.strip()]
        if len(populated) < 2:
            continue
        strings = sum(_numeric_value(value) is None for value in populated)
        unique = len({value.casefold() for value in populated})
        following = rows[position + 1 : position + 4]
        aligned = sum(
            1
            for _next_index, next_values in following
            if sum(bool(value.strip()) for value in next_values[: len(values)]) >= 2
        )
        header_words = sum(
            bool(re.search(r"[A-Za-z]", value)) and len(value) <= 80 for value in populated
        )
        header_labels = sum(
            bool(
                re.fullmatch(
                    r"(?:no\.?|item|product|part number|description|product description|"
                    r"quantity|qty\.?|unit|unit price(?: usd)?|total|lead time|"
                    r"proposal|cisco propose:?|cisco note|notes?)",
                    value.strip(),
                    re.IGNORECASE,
                )
            )
            for value in populated
        )
        numeric_cells = sum(_numeric_value(value) is not None for value in populated)
        long_cells = sum(len(value) > 100 for value in populated)
        score = (
            len(populated) * 2
            + strings
            + unique * 0.5
            + aligned * 2
            + header_words
            + header_labels * 4
            - numeric_cells * 2
            - long_cells * 4
            - position * 0.1
        )
        if score > best_score:
            best_position = position
            best_score = score
    return best_position


def _spreadsheet_layout(
    rows: list[tuple[int, list[str]]],
) -> tuple[list[str], int]:
    """Return composite column labels and the first data-row position.

    Excel exports commonly place a title or a secondary header immediately below
    the main header. Preserve those labels without treating a real line item as a
    header. Key/value sizing sheets receive stable Metric/Value labels instead of
    promoting their first fact (for example Model=Qwen) into a column name.
    """
    if not rows:
        return [], 0
    header_position = _header_row_position(rows)
    header = list(rows[header_position][1])
    populated_pairs = [
        (position, values)
        for position, (_row_index, values) in enumerate(rows)
        if sum(bool(value.strip()) for value in values) >= 2
    ]
    first_header = next((value.strip() for value in header if value.strip()), "")
    recognized_header = bool(
        re.fullmatch(
            r"(?:parameter|metric|name|no\.?|item|product|part number|description|"
            r"product description|component|komponen|category|kategory|line number)",
            first_header,
            re.IGNORECASE,
        )
    )
    label_value_rows = sum(
        bool(values and re.search(r"[A-Za-z]", values[0]))
        and 2 <= sum(bool(value.strip()) for value in values) <= 3
        for _position, values in populated_pairs
    )
    if not recognized_header and label_value_rows >= 6:
        width = max((len(values) for _position, values in populated_pairs), default=2)
        synthetic = ["Metric", "Value", *[f"Value {index}" for index in range(2, width)]]
        return synthetic[:width], populated_pairs[0][0]

    data_position = header_position + 1
    for _ in range(2):
        if data_position >= len(rows):
            break
        values = rows[data_position][1]
        populated = [
            (index, value.strip()) for index, value in enumerate(values) if value.strip()
        ]
        if not populated:
            data_position += 1
            continue
        text = " ".join(value for _index, value in populated)
        if (
            re.search(r"\bgroup name\b|\bsubtotal\b", text, re.IGNORECASE)
            or any(value.startswith("=") for _index, value in populated)
            or re.fullmatch(r"\d+(?:\.\d+)*", populated[0][1])
        ):
            break
        changed = False
        if len(populated) == 1:
            index, value = populated[0]
            if index < len(header) and header[index].strip():
                header[index] = f"{header[index].strip()} > {value}"
            else:
                if index >= len(header):
                    header.extend([""] * (index + 1 - len(header)))
                header[index] = value
            changed = True
        else:
            for index, value in populated:
                if index >= len(header):
                    header.extend([""] * (index + 1 - len(header)))
                if not header[index].strip():
                    header[index] = value
                    changed = True
        if not changed:
            break
        data_position += 1
    return header, data_position


def _range_label(prefix: str, start: int, end: int) -> str:
    if start <= 0:
        return ""
    return f"{prefix} {start}" if start == end else f"{prefix}s {start}-{end}"


def _location(extension: str, chunk: DocumentChunk) -> str:
    if extension == ".pdf":
        return _range_label("Page", chunk.page_start, chunk.page_end) or "PDF source"
    if extension == ".pptx":
        return _range_label("Slide", chunk.page_start, chunk.page_end) or "Presentation source"
    if extension == ".xlsx":
        parts = [f"Sheet {chunk.sheet_name}" if chunk.sheet_name else "", chunk.cell_range]
        return ", ".join(part for part in parts if part) or "Spreadsheet source"
    if extension == ".docx":
        paragraph = _range_label("Paragraph", chunk.paragraph_start, chunk.paragraph_end)
        return ", ".join(part for part in [chunk.heading, paragraph] if part) or "Document source"
    if extension == ".csv":
        return _range_label("Row", chunk.row_start, chunk.row_end) or "CSV source"
    if extension in {".png", ".jpg", ".jpeg"}:
        return "Image text"
    section = chunk.heading
    lines = _range_label("Line", chunk.line_start, chunk.line_end)
    return ", ".join(part for part in [section, lines] if part) or "Text source"


def _merge_chunks(left: DocumentChunk, right: DocumentChunk) -> DocumentChunk:
    return DocumentChunk(
        text=f"{left.text.rstrip()}\n\n{right.text.lstrip()}",
        heading=right.heading or left.heading,
        page_start=min(value for value in [left.page_start, right.page_start] if value > 0)
        if left.page_start or right.page_start
        else 0,
        page_end=max(left.page_end, right.page_end),
        paragraph_start=min(
            value for value in [left.paragraph_start, right.paragraph_start] if value > 0
        )
        if left.paragraph_start or right.paragraph_start
        else 0,
        paragraph_end=max(left.paragraph_end, right.paragraph_end),
        sheet_name=right.sheet_name or left.sheet_name,
        cell_range=right.cell_range or left.cell_range,
        row_start=min(value for value in [left.row_start, right.row_start] if value > 0)
        if left.row_start or right.row_start
        else 0,
        row_end=max(left.row_end, right.row_end),
        line_start=min(value for value in [left.line_start, right.line_start] if value > 0)
        if left.line_start or right.line_start
        else 0,
        line_end=max(left.line_end, right.line_end),
        metadata={**left.metadata, **right.metadata},
    )


def _overlap_key(extension: str, chunk: DocumentChunk) -> tuple[Any, ...]:
    if extension in {'.pdf', '.pptx'}:
        return (chunk.page_start, chunk.page_end, chunk.heading)
    if extension == '.docx':
        return (chunk.paragraph_start, chunk.paragraph_end, chunk.heading)
    if extension == '.xlsx':
        return (chunk.sheet_name, chunk.cell_range)
    if extension == '.csv':
        return (chunk.row_start, chunk.row_end)
    return (chunk.line_start, chunk.line_end, chunk.heading)


class DocumentExtractor:
    def __init__(self) -> None:
        self.encoding = tiktoken.get_encoding("cl100k_base")

    def _size(self, text: str, splitter: str) -> int:
        return len(text) if splitter == "character" else len(self.encoding.encode(text))

    def _tail(self, text: str, amount: int, splitter: str) -> str:
        if amount <= 0:
            return ""
        if splitter == "character":
            return text[-amount:]
        # Preserve semantic boundaries in overlap text. Arbitrary token tails often
        # start halfway through a sentence, which later looks like a factual fragment.
        units = [line.strip() for line in text.splitlines() if line.strip()]
        selected: list[str] = []
        for unit in reversed(units):
            candidate = "\n".join([unit, *selected])
            if self._size(candidate, splitter) > amount:
                if selected:
                    break
                sentences = [
                    sentence.strip()
                    for sentence in re.split(r"(?<=[.!?])\s+(?=[A-Z0-9])", unit)
                    if sentence.strip()
                ]
                for sentence in reversed(sentences):
                    sentence_candidate = " ".join([sentence, *selected])
                    if self._size(sentence_candidate, splitter) > amount:
                        break
                    selected.insert(0, sentence)
                break
            selected.insert(0, unit)
        return "\n".join(selected)

    def _split_text(self, text: str, limit: int, splitter: str) -> list[str]:
        if self._size(text, splitter) <= limit:
            return [text]
        if splitter == "character":
            return [text[index : index + limit] for index in range(0, len(text), limit)]
        units: list[str] = []
        for line in text.splitlines():
            cleaned = line.strip()
            if not cleaned:
                continue
            if self._size(cleaned, splitter) <= limit:
                units.append(cleaned)
                continue
            sentences = [
                sentence.strip()
                for sentence in re.split(r"(?<=[.!?])\s+(?=[A-Z0-9])", cleaned)
                if sentence.strip()
            ]
            for sentence in sentences:
                if self._size(sentence, splitter) <= limit:
                    units.append(sentence)
                    continue
                tokens = self.encoding.encode(sentence)
                units.extend(
                    self.encoding.decode(tokens[index : index + limit]).strip()
                    for index in range(0, len(tokens), limit)
                )

        parts: list[str] = []
        current: list[str] = []
        for unit in units:
            candidate = "\n".join([*current, unit])
            if current and self._size(candidate, splitter) > limit:
                parts.append("\n".join(current))
                current = []
            current.append(unit)
        if current:
            parts.append("\n".join(current))
        return parts

    def _finalize(
        self,
        chunks: list[DocumentChunk],
        file_name: str,
        extension: str,
        settings: dict[str, Any],
    ) -> list[DocumentChunk]:
        splitter = str(settings.get("textSplitter", "token"))
        limit = max(100, int(settings.get("chunkSize", 1000)))
        minimum = max(0, min(limit, int(settings.get("minimumChunkSize", 0))))
        overlap = max(0, min(limit // 2, int(settings.get("chunkOverlap", 0))))

        split_chunks: list[DocumentChunk] = []
        for chunk in chunks:
            parts = self._split_text(chunk.text, limit, splitter)
            for part in parts:
                split_chunks.append(DocumentChunk(**{**chunk.__dict__, "text": part}))

        merged: list[DocumentChunk] = []
        for chunk in split_chunks:
            if (
                merged
                and self._size(chunk.text, splitter) < minimum
                and _overlap_key(extension, merged[-1]) == _overlap_key(extension, chunk)
                and self._size(f"{merged[-1].text}\n\n{chunk.text}", splitter) <= limit
            ):
                merged[-1] = _merge_chunks(merged[-1], chunk)
            else:
                merged.append(chunk)

        previous: DocumentChunk | None = None
        for chunk in merged:
            tail = (
                self._tail(previous.text, overlap, splitter)
                if previous is not None
                and _overlap_key(extension, previous) == _overlap_key(extension, chunk)
                else ""
            )
            if tail and not chunk.text.startswith(tail):
                chunk.text = f"{tail}\n\n{chunk.text}"
            chunk.location = _location(extension, chunk)
            context = []
            if settings.get("enrichedMetadata", True):
                context.append(f"File: {file_name}")
                if settings.get("preserveHeadings", True) and chunk.heading:
                    context.append(f"Section: {chunk.heading}")
                context.append(f"Location: {chunk.location}")
            context.append(chunk.text)
            chunk.embedding_text = "\n".join(context)
            previous = chunk
        return merged

    def extract(
        self, file_path: str, relative_path: str, settings: dict[str, Any]
    ) -> list[DocumentChunk]:
        path = Path(file_path)
        extension = path.suffix.lower()
        if extension == ".pdf":
            chunks = self._extract_pdf(path, settings)
        elif extension == ".xlsx":
            chunks = self._extract_xlsx(path, settings)
        elif extension == ".csv":
            chunks = self._extract_csv(path, settings)
        elif extension == ".pptx":
            chunks = self._extract_pptx(path)
        elif extension in {".txt", ".md"}:
            chunks = self._extract_text(path, settings)
        elif extension in {".png", ".jpg", ".jpeg"}:
            chunks = self._extract_image_ocr(path)
        else:
            chunks = self._extract_docling(path, extension, settings)
        return self._finalize(chunks, relative_path, extension, settings)

    def _shape_text(self, shape: Any) -> list[str]:
        """Extract readable text without allowing a shape group to lose its children."""
        if shape.shape_type == MSO_SHAPE_TYPE.GROUP:
            lines: list[str] = []
            for child in sorted(shape.shapes, key=lambda item: (item.top, item.left)):
                lines.extend(self._shape_text(child))
            return lines
        if getattr(shape, "has_table", False):
            rows = []
            for row in shape.table.rows:
                values = [re.sub(r"\s+", " ", cell.text).strip() for cell in row.cells]
                if any(values):
                    rows.append(" | ".join(values))
            return rows
        if not getattr(shape, "has_text_frame", False):
            return []
        lines = []
        for paragraph in shape.text_frame.paragraphs:
            text = re.sub(r"\s+", " ", paragraph.text).strip()
            if text:
                lines.append(text)
        return lines

    def _extract_pptx(self, path: Path) -> list[DocumentChunk]:
        """Keep every slide as an independent citation and retrieval boundary."""
        presentation = Presentation(path)
        chunks: list[DocumentChunk] = []
        for index, slide in enumerate(presentation.slides, start=1):
            title_shape = slide.shapes.title
            title = re.sub(r"\s+", " ", title_shape.text).strip() if title_shape else ""
            ordered = sorted(
                slide.shapes,
                key=lambda shape: (
                    0 if title_shape is not None and shape == title_shape else 1,
                    shape.top,
                    shape.left,
                ),
            )
            lines: list[str] = []
            seen: set[str] = set()
            for shape in ordered:
                for line in self._shape_text(shape):
                    key = re.sub(r"\s+", " ", line).casefold()
                    if not key or key in seen:
                        continue
                    seen.add(key)
                    lines.append(line)
            if not title:
                title = next(
                    (
                        line
                        for line in lines
                        if 4 <= len(line) <= 120
                        and len(line.split()) <= 14
                        and not re.match(r"^(?:design highlights?|reference|\d+\s*x\b)", line, re.I)
                        and not re.search(r"[.!?]$", line)
                    ),
                    "",
                )
            text = "\n".join(lines).strip()
            if text:
                chunks.append(
                    DocumentChunk(
                        text=text,
                        heading=title,
                        page_start=index,
                        page_end=index,
                        metadata={"slide": index, "title": title},
                    )
                )
        return chunks

    def extract_spreadsheet_cells(self, file_path: str) -> list[SpreadsheetCell]:
        path = Path(file_path)
        if path.suffix.lower() == ".xlsx":
            return self._extract_xlsx_cells(path)
        if path.suffix.lower() == ".csv":
            return self._extract_csv_cells(path)
        return []

    def document_metadata(
        self, relative_path: str, chunks: list[DocumentChunk]
    ) -> dict[str, Any]:
        path = Path(relative_path)
        headings = list(dict.fromkeys(chunk.heading.strip() for chunk in chunks if chunk.heading))
        sections: list[dict[str, str]] = []
        for chunk in chunks:
            text = re.sub(r"\s+", " ", chunk.text).strip()
            if not text:
                continue
            summary = text[:700].rsplit(" ", 1)[0] if len(text) > 700 else text
            sections.append(
                {
                    "heading": chunk.heading,
                    "location": chunk.location,
                    "summary": summary,
                }
            )
        entity_source = " ".join([path.stem, *headings])
        entities = list(
            dict.fromkeys(
                match.group(0)
                for match in re.finditer(
                    r"\b(?:[A-Z][A-Z0-9-]{1,}|[A-Z][a-z]+(?:\s+[A-Z][A-Za-z0-9-]+){1,3})\b",
                    entity_source,
                )
            )
        )[:80]
        overview = " ".join(section["summary"] for section in sections[:4])[:1800]
        return {
            "title": path.stem.replace("_", " ").strip(),
            "extension": path.suffix.lower(),
            "pathSegments": [part for part in path.parts if part not in {".", ".."}],
            "headings": headings[:120],
            "entities": entities,
            "summary": overview,
            "sections": sections[:160],
        }

    def _extract_xlsx_cells(self, path: Path) -> list[SpreadsheetCell]:
        values_book = load_workbook(path, read_only=True, data_only=True)
        formulas_book = load_workbook(path, read_only=True, data_only=False)
        cells: list[SpreadsheetCell] = []
        try:
            for values_sheet, formulas_sheet in zip(
                values_book.worksheets, formulas_book.worksheets, strict=True
            ):
                paired_rows = [
                    (row_index, value_row, formula_row)
                    for row_index, (value_row, formula_row) in enumerate(
                        zip(values_sheet.iter_rows(), formulas_sheet.iter_rows(), strict=False),
                        start=1,
                    )
                    if any(_display_value(cell.value).strip() for cell in value_row)
                ]
                display_rows = [
                    (row_index, [_display_value(cell.value).strip() for cell in value_row])
                    for row_index, value_row, _formula_row in paired_rows
                ]
                headers, data_position = _spreadsheet_layout(display_rows)
                for row_index, value_row, formula_row in paired_rows[data_position:]:
                    displays = [_display_value(cell.value).strip() for cell in value_row]
                    if not any(displays):
                        continue
                    for column_index, (value_cell, formula_cell) in enumerate(
                        zip(value_row, formula_row, strict=False), start=1
                    ):
                        display = _display_value(value_cell.value).strip()
                        formula = (
                            str(formula_cell.value)
                            if formula_cell.data_type == "f" and formula_cell.value is not None
                            else ""
                        )
                        if not display and not formula:
                            continue
                        header = (
                            headers[column_index - 1].strip()
                            if column_index <= len(headers) and headers[column_index - 1].strip()
                            else f"Column {get_column_letter(column_index)}"
                        )
                        raw = _display_value(formula_cell.value).strip()
                        cells.append(
                            SpreadsheetCell(
                                sheet_name=values_sheet.title,
                                row_index=row_index,
                                column_index=column_index,
                                column_name=header,
                                cell_address=f"{get_column_letter(column_index)}{row_index}",
                                display_value=display or raw,
                                raw_value=raw,
                                value_type=str(value_cell.data_type or ""),
                                numeric_value=_numeric_value(value_cell.value),
                                unit=_unit(header, display or raw),
                                formula=formula,
                            )
                        )
        finally:
            values_book.close()
            formulas_book.close()
        return cells

    def _extract_csv_cells(self, path: Path) -> list[SpreadsheetCell]:
        with path.open("r", encoding="utf-8-sig", errors="replace", newline="") as handle:
            rows = csv.reader(handle)
            headers = next(rows, None)
            if not headers:
                return []
            cells: list[SpreadsheetCell] = []
            for row_index, row in enumerate(rows, start=2):
                for column_index, raw in enumerate(row, start=1):
                    display = str(raw).strip()
                    if not display:
                        continue
                    header = (
                        str(headers[column_index - 1]).strip()
                        if column_index <= len(headers) and str(headers[column_index - 1]).strip()
                        else f"Column {get_column_letter(column_index)}"
                    )
                    number = _numeric_value(display)
                    cells.append(
                        SpreadsheetCell(
                            sheet_name="",
                            row_index=row_index,
                            column_index=column_index,
                            column_name=header,
                            cell_address=f"{get_column_letter(column_index)}{row_index}",
                            display_value=display,
                            raw_value=display,
                            value_type="number" if number is not None else "string",
                            numeric_value=number,
                            unit=_unit(header, display),
                        )
                    )
            return cells

    def _extract_pdf(
        self, path: Path, settings: dict[str, Any]
    ) -> list[DocumentChunk]:
        chunks: list[DocumentChunk] = []
        document = pdfium.PdfDocument(path)
        try:
            for index in range(len(document)):
                page = document[index]
                text_page = page.get_textpage()
                try:
                    text = text_page.get_text_range().strip()
                finally:
                    text_page.close()
                    page.close()
                if text:
                    lines = [line.strip() for line in text.splitlines() if line.strip()]
                    reflowed: list[str] = []
                    for line in lines:
                        if (
                            reflowed
                            and not re.search(r"[.!?:;]$", reflowed[-1])
                            and (len(reflowed[-1]) >= 55 or line[:1].islower())
                        ):
                            reflowed[-1] = f"{reflowed[-1]} {line}"
                        else:
                            reflowed.append(line)
                    text = "\n".join(reflowed)
                    chunks.append(
                        DocumentChunk(text=text, page_start=index + 1, page_end=index + 1)
                    )
        finally:
            document.close()

        if chunks:
            return chunks
        if settings.get("pdfOcr", False):
            return self._extract_pdf_ocr(path)
        raise RuntimeError(
            "This PDF contains no readable text. Turn on PDF text recognition and sync again."
        )

    def _extract_pdf_ocr(self, path: Path) -> list[DocumentChunk]:
        """Read image-only PDFs locally without Docling's layout-model download."""
        try:
            from rapidocr import EngineType, RapidOCR
        except ImportError as error:
            raise RuntimeError(
                "PDF text recognition is not available yet. Connect to the internet once and sync again."
            ) from error

        try:
            ocr = RapidOCR(
                params={
                    "Det.engine_type": EngineType.TORCH,
                    "Cls.engine_type": EngineType.TORCH,
                    "Rec.engine_type": EngineType.TORCH,
                }
            )
            document = pdfium.PdfDocument(path)
            chunks: list[DocumentChunk] = []
            try:
                for index in range(len(document)):
                    page = document[index]
                    bitmap = page.render(scale=2)
                    try:
                        result = ocr(bitmap.to_pil())
                    finally:
                        bitmap.close()
                        page.close()
                    text = "\n".join(result.txts or []).strip()
                    if text:
                        chunks.append(
                            DocumentChunk(text=text, page_start=index + 1, page_end=index + 1)
                        )
            finally:
                document.close()
        except Exception as error:
            raise RuntimeError(
                "RavenDock could not download or run local PDF text recognition. "
                "Check your internet connection, then sync again."
            ) from error

        if chunks:
            return chunks
        raise RuntimeError(
            "This scanned PDF could not be read. Try a clearer copy or turn off PDF text recognition to skip it."
        )

    def _extract_image_ocr(self, path: Path) -> list[DocumentChunk]:
        """Index text from standalone images locally."""
        try:
            from PIL import Image
            from rapidocr import EngineType, RapidOCR
        except ImportError as error:
            raise RuntimeError(
                "Image text recognition is not installed yet. Connect once, then sync again."
            ) from error
        try:
            ocr = RapidOCR(
                params={
                    "Det.engine_type": EngineType.TORCH,
                    "Cls.engine_type": EngineType.TORCH,
                    "Rec.engine_type": EngineType.TORCH,
                }
            )
            with Image.open(path) as image:
                result = ocr(image.convert("RGB"))
            text = "\n".join(result.txts or []).strip()
        except Exception as error:
            raise RuntimeError("This image could not be read by local text recognition.") from error
        if not text:
            raise RuntimeError("This image does not contain readable text.")
        return [DocumentChunk(text=text)]

    def _extract_docling(
        self, path: Path, extension: str, settings: dict[str, Any]
    ) -> list[DocumentChunk]:
        pdf_options = PdfPipelineOptions()
        pdf_options.do_ocr = bool(settings.get("pdfOcr", False))
        converter = DocumentConverter(
            format_options={InputFormat.PDF: PdfFormatOption(pipeline_options=pdf_options)}
        )
        document = converter.convert(path).document
        tokenizer = OpenAITokenizer(
            tokenizer=self.encoding,
            max_tokens=max(100, int(settings.get("chunkSize", 1000))),
        )
        chunker = HybridChunker(tokenizer=tokenizer, merge_peers=True, repeat_table_header=True)
        chunks: list[DocumentChunk] = []
        for item in chunker.chunk(document):
            metadata = item.meta.model_dump(mode="json")
            doc_items = metadata.get("doc_items") or []
            pages = sorted(
                {
                    int(prov["page_no"])
                    for doc_item in doc_items
                    for prov in doc_item.get("prov", [])
                    if isinstance(prov.get("page_no"), int)
                }
            )
            paragraphs = sorted(
                {
                    int(match.group(1)) + 1
                    for doc_item in doc_items
                    if (match := re.fullmatch(r"#/texts/(\d+)", str(doc_item.get("self_ref", ""))))
                }
            )
            headings = metadata.get("headings") or []
            chunks.append(
                DocumentChunk(
                    text=item.text.strip(),
                    heading=" > ".join(str(value) for value in headings),
                    page_start=pages[0] if pages else 0,
                    page_end=pages[-1] if pages else 0,
                    paragraph_start=paragraphs[0] if paragraphs else 0,
                    paragraph_end=paragraphs[-1] if paragraphs else 0,
                    metadata=metadata,
                )
            )
        return [chunk for chunk in chunks if chunk.text]

    def _table_chunks(
        self,
        rows: Iterable[tuple[int, list[str]]],
        header: list[str],
        settings: dict[str, Any],
        sheet_name: str = "",
    ) -> list[DocumentChunk]:
        limit = max(100, int(settings.get("chunkSize", 1000)))
        splitter = str(settings.get("textSplitter", "token"))
        chunks: list[DocumentChunk] = []
        current: list[tuple[int, list[str]]] = []

        def render(values: list[tuple[int, list[str]]]) -> str:
            width = max([len(header), *(len(row) for _, row in values)], default=len(header))
            normalized_header = (header + [""] * width)[:width]
            lines = [
                "| " + " | ".join(normalized_header) + " |",
                "| " + " | ".join(["---"] * width) + " |",
            ]
            lines.extend(
                "| " + " | ".join((row + [""] * width)[:width]) + " |" for _, row in values
            )
            return "\n".join(lines)

        def flush() -> None:
            if not current:
                return
            start, end = current[0][0], current[-1][0]
            max_column = max([len(header), *(len(row) for _, row in current)], default=1)
            chunks.append(
                DocumentChunk(
                    text=render(current),
                    sheet_name=sheet_name,
                    cell_range=f"A{start}:{get_column_letter(max_column)}{end}",
                    row_start=start,
                    row_end=end,
                )
            )
            current.clear()

        for numbered_row in rows:
            candidate = [*current, numbered_row]
            if current and self._size(render(candidate), splitter) > limit:
                flush()
            current.append(numbered_row)
        flush()
        return chunks

    def _extract_xlsx(self, path: Path, settings: dict[str, Any]) -> list[DocumentChunk]:
        # Cached values make spreadsheet calculations searchable instead of exposing formulas.
        workbook = load_workbook(path, read_only=True, data_only=True)
        chunks: list[DocumentChunk] = []
        try:
            for sheet in workbook.worksheets:
                def populated_rows() -> Iterable[tuple[int, list[str]]]:
                    for index, values in enumerate(
                        sheet.iter_rows(values_only=True), start=1
                    ):
                        row = [_display_value(value) for value in values]
                        while row and not row[-1]:
                            row.pop()
                        if row and any(cell.strip() for cell in row):
                            yield index, row

                rows = list(populated_rows())
                if not rows:
                    continue
                header, data_position = _spreadsheet_layout(rows)
                data = rows[data_position:]
                if not data:
                    data = [rows[_header_row_position(rows)]]
                chunks.extend(self._table_chunks(data, header, settings, sheet.title))
        finally:
            workbook.close()
        return chunks

    def _extract_csv(self, path: Path, settings: dict[str, Any]) -> list[DocumentChunk]:
        with path.open("r", encoding="utf-8-sig", errors="replace", newline="") as handle:
            rows = iter(csv.reader(handle))
            first = next(rows, None)
            if first is None:
                return []
            header = [str(value) for value in first]
            second = next(rows, None)
            numbered: Iterable[tuple[int, list[str]]]
            if second is None:
                numbered = [(1, header)]
            else:
                numbered = chain(
                    [(2, [str(value) for value in second])],
                    (
                        (index, [str(value) for value in row])
                        for index, row in enumerate(rows, 3)
                    ),
                )
            return self._table_chunks(numbered, header, settings)

    def _extract_text(self, path: Path, settings: dict[str, Any]) -> list[DocumentChunk]:
        lines = path.read_text(encoding="utf-8", errors="replace").splitlines()
        chunks: list[DocumentChunk] = []
        heading = ""
        start = 1
        buffer: list[str] = []
        limit = max(100, int(settings.get("chunkSize", 1000)))
        splitter = str(settings.get("textSplitter", "token"))
        for number, line in enumerate(lines, start=1):
            if path.suffix.lower() == ".md" and line.lstrip().startswith("#"):
                if buffer:
                    chunks.append(
                        DocumentChunk(
                            text="\n".join(buffer),
                            heading=heading,
                            line_start=start,
                            line_end=number - 1,
                        )
                    )
                    buffer = []
                    start = number
                heading = line.lstrip("#").strip()
            candidate = "\n".join([*buffer, line])
            if buffer and self._size(candidate, splitter) > limit:
                chunks.append(
                    DocumentChunk(
                        text="\n".join(buffer),
                        heading=heading,
                        line_start=start,
                        line_end=number - 1,
                    )
                )
                start = number
                buffer = []
            buffer.append(line)
        if buffer:
            chunks.append(
                DocumentChunk(
                    text="\n".join(buffer),
                    heading=heading,
                    line_start=start,
                    line_end=len(lines),
                )
            )
        return chunks
